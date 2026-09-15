from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
PPT = OUT / "VGGT_TRELLIS_Affostruction架构讲解_4页.pptx"
NOTES = OUT / "VGGT_TRELLIS_Affostruction架构讲解_4页_演讲稿.md"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
BG = RGBColor(8, 17, 30)
PANEL = RGBColor(18, 31, 50)
PANEL2 = RGBColor(24, 40, 62)
WHITE = RGBColor(240, 245, 250)
MUTED = RGBColor(154, 171, 190)
CYAN = RGBColor(54, 207, 220)
BLUE = RGBColor(74, 139, 255)
GREEN = RGBColor(58, 203, 132)
PURPLE = RGBColor(167, 112, 255)
ORANGE = RGBColor(245, 166, 54)


def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=14, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = Inches(0.06); frame.margin_right = Inches(0.06)
    frame.margin_top = Inches(0.03); frame.margin_bottom = Inches(0.03)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return shape


def bullets(slide, values, x, y, w, h, size=12.5, color=WHITE, gap=7):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = Inches(0.08); frame.margin_right = Inches(0.04)
    for i, value in enumerate(values):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = "• " + value; p.font.name = FONT; p.font.size = Pt(size)
        p.font.color.rgb = color; p.space_after = Pt(gap)
    return shape


def page_title(slide, n, heading, subtitle):
    text(slide, f"0{n}", 0.44, 0.26, 0.55, 0.42, 13, CYAN, True)
    text(slide, heading, 1.02, 0.20, 11.75, 0.58, 25, WHITE, True)
    text(slide, subtitle, 1.04, 0.77, 11.55, 0.30, 10.5, MUTED)
    rect(slide, 0.44, 1.12, 12.35, 0.025, CYAN, radius=False)


def footer(slide):
    text(slide, "VGGT × TRELLIS × Affostruction · Sparse-view Image-to-3D", 0.47, 7.15, 7.5, 0.19, 8.7, MUTED)


def node(slide, heading, detail, x, y, w, color, h=1.12):
    rect(slide, x, y, w, h, PANEL2, line=color)
    text(slide, heading, x + 0.09, y + 0.09, w - 0.18, 0.32, 14, color, True, PP_ALIGN.CENTER)
    text(slide, detail, x + 0.09, y + 0.43, w - 0.18, h - 0.50, 10.1, WHITE, False, PP_ALIGN.CENTER)


def arrow(slide, x, y, w=0.32, color=CYAN):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.30))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.color.rgb = color


def notes(slide, value):
    frame = slide.notes_slide.notes_text_frame
    frame.clear(); frame.paragraphs[0].text = value


prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]
scripts = []


# Slide 1
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 1, "总体思路：把图像条件变成三维空间条件", "吸收 Affostruction 的空间条件化生成与稀疏流思想，适配 VGGT + TRELLIS 两阶段重建")

node(s, "VGGT", "相机 K/T、depth、point map、confidence", 0.55, 1.50, 2.05, CYAN)
arrow(s, 2.72, 1.90, 0.34)
node(s, "SS 条件流", "Pixel → 16³ voxel\n生成 active support", 3.18, 1.50, 2.13, GREEN)
arrow(s, 5.43, 1.90, 0.34)
node(s, "SLat 残差流", "64³ voxel → Pixel\n生成局部属性修正", 5.89, 1.50, 2.20, PURPLE)
arrow(s, 8.21, 1.90, 0.34)
node(s, "TRELLIS 先验", "遮挡/未观测区域\n保持生成完整性", 8.67, 1.50, 2.08, BLUE)
arrow(s, 10.87, 1.90, 0.34)
node(s, "3D Asset", "Gaussian / mesh\n新视角渲染", 11.33, 1.50, 1.45, ORANGE)

rect(s, 0.55, 3.04, 3.86, 3.13, PANEL, line=CYAN)
text(s, "从 Affostruction 直接采用", 0.82, 3.30, 3.32, 0.38, 17, CYAN, True)
bullets(s, [
    "RGBD 像素反投影并附着到三维 voxel",
    "空间条件 token + 3D positional encoding",
    "Transformer cross-attention 条件注入",
    "官方 SS 768×12 Conditional Flow Matching",
], 0.86, 3.88, 3.15, 1.85, 12.6)

rect(s, 4.70, 3.04, 3.89, 3.13, PANEL, line=PURPLE)
text(s, "针对图像重建进行适配", 4.97, 3.30, 3.35, 0.38, 17, PURPLE, True)
bullets(s, [
    "不使用文本、CLIP 或 affordance prompt",
    "条件改为 VGGT 几何 + DINO 语义 + RGB 纹理",
    "SLat active voxel 回投所有输入视图",
    "多视图证据按可见性与一致性融合",
], 5.01, 3.88, 3.15, 1.85, 12.6)

rect(s, 8.88, 3.04, 3.89, 3.13, PANEL, line=GREEN)
text(s, "总体控制原则", 9.15, 3.30, 3.35, 0.38, 17, GREEN, True)
bullets(s, [
    "SS 决定“哪里存在结构”",
    "SLat 决定“结构点具有什么属性”",
    "观测充分区域由 VGGT 校正",
    "证据不足区域回归 TRELLIS 生成先验",
], 9.19, 3.88, 3.13, 1.85, 12.6)
footer(s)
n = """当前架构吸收 Affostruction 的核心并不是增加一个普通 Adapter，而是改变条件进入生成模型的方式。VGGT 输出相机、深度、点图和置信度，SS 阶段把二维像素证据转换成十六立方三维条件，生成 active support；SLat 阶段从六十四立方 active voxel 再投影回图像，读取对应的语义与纹理。Affostruction 的空间 token、三维位置编码、cross-attention 和 Conditional Flow Matching 被保留，但文本或 affordance 条件没有被采用。系统最终通过置信度决定由观测证据还是 TRELLIS 先验控制某个区域。"""
notes(s, n); scripts.append((1, "总体思路", n))


# Slide 2
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 2, "SS：RGBD 体素条件驱动的官方生成流", "将 VGGT 的确定性几何变成 Affostruction SS Flow 可读取的紧凑三维 token")
node(s, "RGB / Mask", "按 frame ID 精确对齐\n缺失视图不重排", 0.55, 1.48, 1.65, CYAN)
arrow(s, 2.30, 1.88)
node(s, "DINO ViT-L/14", "冻结编码器\n1024-D patch map", 2.70, 1.48, 1.75, BLUE)
arrow(s, 4.55, 1.88)
node(s, "RGBD 反投影", "K、T、depth\nPixel → canonical 3D", 4.95, 1.48, 1.83, GREEN)
arrow(s, 6.88, 1.88)
node(s, "16³ 融合", "voxel 内采样\n跨有效视图平均", 7.28, 1.48, 1.70, CYAN)
arrow(s, 9.08, 1.88)
node(s, "条件 Token", "LayerNorm + 3D PE\ncompact + valid mask", 9.48, 1.48, 1.70, BLUE)
arrow(s, 11.28, 1.88)
node(s, "SS Flow", "768×12\nvelocity field", 11.68, 1.48, 1.10, GREEN)

rect(s, 0.55, 3.04, 4.03, 3.13, PANEL, line=BLUE)
text(s, "空间条件构造", 0.82, 3.31, 3.49, 0.37, 17, BLUE, True)
text(s, "Xc=[(u−cx)d/fx, (v−cy)d/fy, d]", 0.83, 3.88, 3.46, 0.38, 13.3, WHITE, True, PP_ALIGN.CENTER)
text(s, "Xw = Tc2w Xc   →   q = floor[(Xw+0.5)·16]", 0.74, 4.31, 3.65, 0.43, 12.5, WHITE, True, PP_ALIGN.CENTER)
bullets(s, [
    "每个 voxel 先求平均采样位置，再用 grid_sample 读取 DINO",
    "仅观测 voxel 进入 condition；view_valid_mask 排除 padding",
], 0.87, 4.94, 3.28, 0.90, 11.7, WHITE, 6)

rect(s, 4.84, 3.04, 3.68, 3.13, PANEL, line=GREEN)
text(s, "官方模型拓扑", 5.11, 3.31, 3.14, 0.37, 17, GREEN, True)
bullets(s, [
    "SparseStructureFlowModel",
    "resolution 16³；in/out 8 channel",
    "hidden 768；12 blocks；12 heads",
    "absolute PE；Q/K RMS norm",
    "DINO 冻结；SS Flow 全参数微调",
], 5.14, 3.91, 3.02, 1.74, 12.2, WHITE, 7)

rect(s, 8.79, 3.04, 3.98, 3.13, PANEL, line=GREEN)
text(s, "Conditional Flow Matching", 9.06, 3.31, 3.44, 0.37, 17, GREEN, True)
text(s, "xt=(1−t)x0+[σmin+(1−σmin)t]ε", 9.00, 3.93, 3.55, 0.42, 13.1, WHITE, True, PP_ALIGN.CENTER)
text(s, "v*=(1−σmin)ε−x0", 9.20, 4.41, 3.15, 0.38, 13.8, WHITE, True, PP_ALIGN.CENTER)
text(s, "LSS = ‖vθ(xt,t,C)−v*‖²", 9.19, 4.88, 3.17, 0.40, 14.1, GREEN, True, PP_ALIGN.CENTER)
text(s, "ODE + CFG → active voxel support", 9.11, 5.51, 3.31, 0.31, 11.7, MUTED, True, PP_ALIGN.CENTER)
footer(s)
n = """SS 的条件构造分为特征提取、三维反投影和跨视图聚合。冻结 DINO 从每个输入视图生成一千零二十四维 patch feature；VGGT 或数据集深度结合相机内外参把像素恢复到 canonical 三维空间，再量化到十六立方 voxel。同一 voxel 中的像素先确定平均二维采样位置，通过 grid_sample 读取 DINO 特征，之后在所有有效视图间平均。聚合结果经过 LayerNorm 并加入三维位置编码，只有被观测到的 voxel 被压缩为条件 token。官方 Affostruction SS Flow 使用十二层、宽度七百六十八的 Transformer，通过 cross-attention 读取这些条件，训练目标为 Conditional Flow Matching 速度 MSE，推理通过 ODE 和 CFG 得到 active support。"""
notes(s, n); scripts.append((2, "SS 条件流", n))


# Slide 3
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 3, "SLat：像素对齐、置信度融合与受控残差", "把 Affostruction 稀疏条件流转化为冻结 TRELLIS 先验上的图像空间校正器")
node(s, "64³ active voxel", "SS support\ncanonical coordinate", 0.55, 1.48, 1.65, GREEN)
arrow(s, 2.29, 1.88)
node(s, "投影所有视图", "Kaolin/OpenCV K,T\n3D → sampling grid", 2.70, 1.48, 1.75, BLUE)
arrow(s, 4.54, 1.88)
node(s, "双频特征", "high: DINO/VGGT\nlow: shallow RGB CNN", 4.95, 1.48, 1.75, PURPLE)
arrow(s, 6.79, 1.88)
node(s, "可信融合", "confidence · visibility\ndepth · angle", 7.20, 1.48, 1.65, CYAN)
arrow(s, 8.94, 1.88)
node(s, "Sparse Flow", "768×12 cross-attn\npredict Δvraw", 9.35, 1.48, 1.55, PURPLE)
arrow(s, 10.99, 1.88)
node(s, "Gated Residual", "frozen base +\nbounded correction", 11.40, 1.48, 1.35, GREEN)

rect(s, 0.55, 3.04, 4.02, 3.13, PANEL, line=CYAN)
text(s, "逐 voxel 多视图证据", 0.82, 3.31, 3.48, 0.37, 17, CYAN, True)
text(s, "αᵢᵛ = cᵢᵛ · mᵢᵛ · wdepth · wangle", 0.82, 3.91, 3.48, 0.42, 14.1, WHITE, True, PP_ALIGN.CENTER)
text(s, "fᵢ = Σᵥ αᵢᵛ[fᴴ⊕fᴸ] / (Σᵥ αᵢᵛ+ε)", 0.76, 4.40, 3.60, 0.42, 13.1, WHITE, True, PP_ALIGN.CENTER)
bullets(s, [
    "mask/occlusion 排除背景与被遮挡观测",
    "depth 与 normal angle 衡量几何一致性",
    "融合后投影为 1024-D condition token",
], 0.89, 5.02, 3.18, 0.89, 11.5, WHITE, 5)

rect(s, 4.84, 3.04, 3.70, 3.13, PANEL, line=PURPLE)
text(s, "Affostruction 稀疏残差流", 5.11, 3.31, 3.16, 0.37, 17, PURPLE, True)
bullets(s, [
    "SLatFlowModel on 64³ sparse support",
    "in/out 8 channel；hidden 768",
    "12 blocks；12 heads；patch size 2",
    "cross-attention 读取像素对齐 condition",
    "训练 spatial flow + conditioner",
], 5.14, 3.91, 3.02, 1.74, 12.0, WHITE, 7)

rect(s, 8.81, 3.04, 3.96, 3.13, PANEL, line=GREEN)
text(s, "冻结 TRELLIS 先验", 9.08, 3.31, 3.42, 0.37, 17, GREEN, True)
text(s, "L = λres · RMS(vbase)", 9.24, 3.93, 3.10, 0.38, 13.8, WHITE, True, PP_ALIGN.CENTER)
text(s, "Δv = c · L · tanh(Δvraw / L)", 9.08, 4.39, 3.41, 0.42, 13.3, WHITE, True, PP_ALIGN.CENTER)
text(s, "vfinal = vfrozen + Δv", 9.32, 4.87, 2.94, 0.41, 14.5, GREEN, True, PP_ALIGN.CENTER)
text(s, "低可信区域 → correction≈0 → 保留生成先验", 9.06, 5.48, 3.45, 0.38, 11.3, MUTED, True, PP_ALIGN.CENTER)
footer(s)
n = """SLat 阶段从 SS 的 active voxel 出发，将每个三维坐标投影到所有有效输入视图，在相同像素位置采样高层 DINO 或 VGGT 特征以及浅层 RGB CNN 特征。高层特征负责部件语义，低层特征保留颜色和局部纹理。每个视图的融合权重由 VGGT confidence、可见性、前景遮挡、深度一致性和法向入射角共同决定。融合后的 token 通过 cross-attention 输入 Affostruction SLatFlowModel。这个稀疏流不替换 TRELLIS，而是预测 residual。TRELLIS 原生 SLat velocity 完全冻结，残差先按 base velocity RMS 确定尺度，再用 tanh 限幅并乘 voxel confidence，因此有证据区域得到校正，弱证据区域保留原始先验。"""
notes(s, n); scripts.append((3, "SLat 可信残差流", n))


# Slide 4
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 4, "训练、推理与参数更新闭环", "两阶段分开优化，统一坐标与条件契约，在生成完整性和重建约束间分配控制权")

node(s, "训练 SS", "RGBD condition\nCFM velocity MSE", 0.55, 1.47, 1.75, GREEN)
arrow(s, 2.41, 1.87)
node(s, "生成 Support", "EMA / CFG / ODE\n16³ → active voxel", 2.84, 1.47, 1.91, GREEN)
arrow(s, 4.86, 1.87)
node(s, "训练 SLat", "frozen base teacher\nlearn target residual", 5.29, 1.47, 1.86, PURPLE)
arrow(s, 7.26, 1.87)
node(s, "联合采样", "base velocity\n+ confidence residual", 7.69, 1.47, 1.94, PURPLE)
arrow(s, 9.74, 1.87)
node(s, "TRELLIS Decode", "Gaussian / mesh\nRGB/mask/depth/geometry", 10.17, 1.47, 2.18, ORANGE)

rect(s, 0.55, 3.02, 3.85, 3.18, PANEL, line=GREEN)
text(s, "SS 参数与目标", 0.82, 3.29, 3.31, 0.37, 17, GREEN, True)
bullets(s, [
    "训练：Affostruction SS Flow 全部参数",
    "冻结：DINO 与 RGBD conditioner",
    "目标：LSS = velocity CFM-MSE",
    "CFG dropout 学习有/无条件速度场",
    "EMA 权重用于稳定生成采样",
], 0.88, 3.90, 3.12, 1.78, 12.2, WHITE, 7)

rect(s, 4.68, 3.02, 4.02, 3.18, PANEL, line=PURPLE)
text(s, "SLat 参数与目标", 4.95, 3.29, 3.48, 0.37, 17, PURPLE, True)
bullets(s, [
    "训练：sparse residual flow + conditioner",
    "冻结：TRELLIS native SLat flow",
    "Lflow：最终 velocity 对齐目标速度",
    "Lendpoint：反推 clean x0；Lprior：低可信归零",
    "Ldecoded：RGB / mask / depth / feature / geometry",
], 5.01, 3.90, 3.30, 1.78, 11.8, WHITE, 6)

rect(s, 8.98, 3.02, 3.79, 3.18, PANEL, line=CYAN)
text(s, "统一工程契约", 9.25, 3.29, 3.25, 0.37, 17, CYAN, True)
bullets(s, [
    "frame ID 同步 RGB、depth、mask、K、T",
    "训练/推理统一 canonical frame 与 camera convention",
    "DDP 多 GPU；BF16 前向；FP32 主权重",
    "gradient checkpointing 控制激活显存",
    "核心原则：观测校正 + 未观测生成",
], 9.31, 3.90, 3.08, 1.78, 11.8, WHITE, 6)

rect(s, 0.55, 6.39, 12.22, 0.45, RGBColor(22, 57, 61), line=CYAN)
text(s, "Pixel → Voxel → Pixel-aligned Evidence → Gated Structured Latent → 3D Asset", 0.75, 6.44, 11.81, 0.31, 14, CYAN, True, PP_ALIGN.CENTER)
footer(s)
n = """训练分成 SS 和 SLat 两个阶段。SS 中只更新官方 Affostruction SS Flow，DINO 和 RGBD conditioner 冻结，使用 CFM 速度 MSE 和 EMA。得到 SS 模型后生成 active support。SLat 中冻结 TRELLIS 原生 SLat flow，把它作为 base teacher，训练独立 residual flow 和 conditioner，使最终速度对齐目标速度，同时使用 endpoint、低置信先验保持和解码资产监督。推理时 SS 先生成 support，SLat 在每个 ODE 步将 frozen base 与 gated residual 相加，再交给 TRELLIS decoder。整个闭环必须保持 frame ID、相机约定、canonical 坐标以及条件 mask 一致。最终方法可以概括为 Pixel 到 Voxel，再从 Voxel 回到像素证据，最后形成受控 Structured Latent。"""
notes(s, n); scripts.append((4, "训练与推理闭环", n))


OUT.mkdir(parents=True, exist_ok=True)
prs.save(PPT)
md = ["# VGGT + TRELLIS + Affostruction 架构讲解——4页演讲稿", ""]
for index, heading, body in scripts:
    md.extend([f"## 第 {index} 页：{heading}", "", body, ""])
NOTES.write_text("\n".join(md), encoding="utf-8")
print(PPT)
print(NOTES)
