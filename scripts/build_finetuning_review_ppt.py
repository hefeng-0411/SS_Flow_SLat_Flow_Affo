from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parents[1] / "reports"
PPT_PATH = OUT_DIR / "当前SS_SLat_Affostruction微调架构.pptx"
SCRIPT_PATH = OUT_DIR / "当前SS_SLat_Affostruction微调架构_演讲稿.md"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
BG = RGBColor(9, 18, 32)
PANEL = RGBColor(18, 31, 50)
PANEL_2 = RGBColor(23, 40, 63)
WHITE = RGBColor(240, 245, 250)
MUTED = RGBColor(154, 171, 190)
CYAN = RGBColor(56, 205, 221)
BLUE = RGBColor(74, 139, 255)
GREEN = RGBColor(59, 203, 133)
ORANGE = RGBColor(245, 166, 55)
RED = RGBColor(242, 91, 104)
PURPLE = RGBColor(166, 113, 255)


def add_rect(slide, x, y, w, h, fill, radius=True, line=None, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    return shape


def add_text(slide, text, x, y, w, h, size=18, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=15, color=WHITE,
                   bullet=False, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.06)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if bullet else "") + line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
    return box


def title(slide, index, headline, sub=None):
    add_text(slide, f"0{index}", 0.45, 0.30, 0.6, 0.42, 13, CYAN, True)
    add_text(slide, headline, 1.05, 0.24, 11.7, 0.58, 26, WHITE, True)
    if sub:
        add_text(slide, sub, 1.07, 0.78, 11.6, 0.32, 10.5, MUTED)
    add_rect(slide, 0.45, 1.14, 12.35, 0.025, CYAN, radius=False)


def footer(slide, text="SS Flow · SLat Flow · VGGT + TRELLIS"):
    add_text(slide, text, 0.48, 7.15, 8.5, 0.20, 9, MUTED)


def arrow(slide, x, y, w=0.38, color=CYAN):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(0.32))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.color.rgb = color
    return shp


def flow_node(slide, heading, detail, x, y, w, color):
    add_rect(slide, x, y, w, 1.12, PANEL_2, line=color)
    add_text(slide, heading, x + 0.12, y + 0.10, w - 0.24, 0.32, 15, color, True, PP_ALIGN.CENTER)
    add_text(slide, detail, x + 0.12, y + 0.44, w - 0.24, 0.53, 10.5, WHITE, False, PP_ALIGN.CENTER)


def metric_row(slide, label, values, y, better="high"):
    xs = [3.25, 5.45, 7.65, 9.85]
    add_text(slide, label, 0.62, y, 2.35, 0.36, 12, WHITE, True)
    nums = [float(v) for v in values]
    best = max(nums) if better == "high" else min(nums)
    for x, v in zip(xs, nums):
        color = GREEN if abs(v - best) < 1e-8 else WHITE
        add_text(slide, f"{v:.3f}", x, y, 1.45, 0.36, 12.5, color, abs(v - best) < 1e-8, PP_ALIGN.CENTER)


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    tf.paragraphs[0].text = text


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

notes = []

# 1 — Cover
s = prs.slides.add_slide(blank)
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = BG
add_rect(s, 0.0, 0.0, 0.14, H, CYAN, radius=False)
add_text(s, "VGGT × TRELLIS × Affostruction", 0.65, 0.62, 8.5, 0.42, 16, CYAN, True)
add_text(s, "当前 SS / SLat 微调架构", 0.65, 1.20, 11.6, 0.88, 34, WHITE, True)
add_text(s, "从稀疏多视图几何条件，到冻结先验的结构化潜变量残差流", 0.68, 2.10, 11.4, 0.42, 17, MUTED)
add_rect(s, 0.68, 3.00, 3.63, 1.33, PANEL, line=BLUE)
add_text(s, "SS", 0.88, 3.19, 0.7, 0.36, 22, BLUE, True)
add_text(s, "官方 768×12 全量流模型\nRGBD voxel 条件 + CFM", 1.55, 3.12, 2.45, 0.75, 13, WHITE)
arrow(s, 4.52, 3.50, 0.48)
add_rect(s, 5.22, 3.00, 4.34, 1.33, PANEL, line=PURPLE)
add_text(s, "SLat", 5.43, 3.19, 1.0, 0.36, 22, PURPLE, True)
add_text(s, "冻结 TRELLIS 先验\n像素对齐稀疏残差流", 6.45, 3.12, 2.72, 0.75, 13, WHITE)
add_rect(s, 0.68, 5.10, 11.58, 0.80, PANEL_2, line=ORANGE)
add_text(s, "核心设计", 0.92, 5.30, 1.15, 0.35, 14, ORANGE, True)
add_text(s, "SS 回归官方条件流；SLat 冻结 TRELLIS 先验，仅学习可信的像素对齐残差", 2.15, 5.24, 9.75, 0.44, 16, WHITE, True)
add_text(s, "项目技术汇报 · 6页精简版", 0.69, 6.55, 4.2, 0.30, 11, MUTED)
note = """本页给出整体架构的核心定位。项目由两个阶段组成：SS 负责生成稀疏结构支撑，SLat 负责在支撑坐标上生成外观和局部几何属性。在 VGGT 加 TRELLIS 基础上，系统加入 Affostruction 官方 SS 全量条件流，以及冻结 TRELLIS 原生 SLat 先验的像素对齐残差流。整个设计的目标是把可观测几何注入模型，同时尽量不破坏 TRELLIS 的生成分布。"""
set_notes(s, note); notes.append((1, "标题与结论", note))

# 2 — End-to-end
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
title(s, 2, "当前微调：端到端最简流程", "输入只使用稀疏多视图图像；SS 定结构，SLat 定属性")
flow_node(s, "稀疏多视图", "RGB / mask / frame ID\nK、T、depth、confidence", 0.55, 1.55, 2.05, CYAN)
arrow(s, 2.72, 1.94)
flow_node(s, "VGGT + DINO", "相机与点图对齐\n高层语义特征", 3.18, 1.55, 1.95, BLUE)
arrow(s, 5.25, 1.94)
flow_node(s, "SS Flow", "16³ voxel condition\nCFM → active support", 5.70, 1.55, 1.95, GREEN)
arrow(s, 7.77, 1.94)
flow_node(s, "SLat Residual", "64³ active voxel\n跨视图像素对齐", 8.22, 1.55, 2.05, PURPLE)
arrow(s, 10.39, 1.94)
flow_node(s, "TRELLIS Decode", "Gaussian / mesh\n新视角渲染", 10.84, 1.55, 1.92, ORANGE)

add_rect(s, 0.55, 3.23, 5.92, 2.92, PANEL)
add_text(s, "SS：学习“哪里有结构”", 0.82, 3.46, 5.30, 0.39, 18, GREEN, True)
add_rich_lines(s, [
    "RGBD 反投影 → 16³ voxel",
    "跨视图平均 DINO token + 3D 位置编码",
    "全量训练官方 Affostruction SS Flow",
    "输出 active voxel support",
], 0.86, 3.98, 5.20, 1.78, 14, WHITE, True, 11)

add_rect(s, 6.72, 3.23, 6.04, 2.92, PANEL)
add_text(s, "SLat：学习“每个结构点是什么”", 6.99, 3.46, 5.50, 0.39, 18, PURPLE, True)
add_rich_lines(s, [
    "投影 active voxel → 每个相机视图",
    "采样高层语义 + 低层 RGB 纹理",
    "置信度 / 遮挡 / 深度 / 角度加权融合",
    "受限残差修正冻结 TRELLIS velocity",
], 7.03, 3.98, 5.28, 1.78, 14, WHITE, True, 11)
footer(s)
note = """端到端流程可以压缩为五步。第一步读取严格按 frame ID 对齐的稀疏多视图数据。第二步由 VGGT 提供相机、深度、点图和置信度，DINO 提供语义特征。第三步 SS 在 16 立方网格上生成 active support。第四步 SLat 把这些 active voxel 投影回所有输入视图，采样语义与纹理特征并融合，然后仅预测相对原始 TRELLIS 的速度残差。第五步由 TRELLIS 解码器生成 Gaussian 和 mesh。这个分工避免让一个模块同时承担拓扑、材质和多视图一致性。"""
set_notes(s, note); notes.append((2, "端到端流程", note))

# 3 — SS algorithm
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
title(s, 3, "SS 微调：RGBD 体素条件驱动的官方生成流", "在 VGGT 几何输出与 TRELLIS 稀疏结构阶段之间建立显式三维条件")
flow_node(s, "多视图 RGBD", "image / mask / depth\nK、camera-to-world", 0.55, 1.55, 2.00, CYAN)
arrow(s, 2.68, 1.94, 0.34)
flow_node(s, "DINO 特征", "冻结 ViT-L/14\n1024-D patch token", 3.14, 1.55, 1.93, BLUE)
arrow(s, 5.20, 1.94, 0.34)
flow_node(s, "三维反投影", "像素 → canonical 3D\n量化到 16³ voxel", 5.66, 1.55, 2.02, GREEN)
arrow(s, 7.81, 1.94, 0.34)
flow_node(s, "条件融合", "同 voxel 跨视图平均\nLayerNorm + 3D PE", 8.27, 1.55, 2.05, CYAN)
arrow(s, 10.45, 1.94, 0.34)
flow_node(s, "SS Flow", "768×12 cross-attn\n预测 latent velocity", 10.91, 1.55, 1.87, GREEN)

add_rect(s, 0.55, 3.21, 4.00, 2.95, PANEL, line=BLUE)
add_text(s, "新增的条件表达", 0.83, 3.46, 3.45, 0.36, 17, BLUE, True)
add_rich_lines(s, [
    "VGGT depth / camera 定义显式三维位置",
    "DINO token 提供局部语义与外观先验",
    "只保留实际观测 voxel，padding 通过 mask 排除",
    "缺失编号图像按 frame ID 对齐，不发生索引漂移",
], 0.88, 3.96, 3.30, 1.77, 12.9, WHITE, True, 9)

add_rect(s, 4.79, 3.21, 3.72, 2.95, PANEL, line=GREEN)
add_text(s, "微调拓扑", 5.07, 3.46, 3.15, 0.36, 17, GREEN, True)
add_rich_lines(s, [
    "Affostruction SparseStructureFlowModel",
    "16³、8 latent channel、hidden 768",
    "12 blocks / 12 heads / absolute PE",
    "SS Flow 全参数训练；DINO 冻结",
], 5.10, 3.96, 3.04, 1.77, 12.9, WHITE, True, 9)

add_rect(s, 8.75, 3.21, 4.02, 2.95, PANEL, line=GREEN)
add_text(s, "生成方法", 9.03, 3.46, 3.46, 0.36, 17, GREEN, True)
add_rich_lines(s, [
    "Conditional Flow Matching 训练速度场",
    "cross-attention 注入紧凑 voxel token",
    "CFG dropout 学习有条件/无条件分支",
    "ODE 采样得到 active voxel support",
], 9.07, 3.96, 3.31, 1.77, 12.9, WHITE, True, 9)
footer(s)
note = """SS 阶段是在 VGGT 几何输出与 TRELLIS 稀疏结构生成之间加入显式三维条件。输入图像先经过冻结 DINO，得到一千零二十四维 patch token；对应深度和相机把像素反投影到 canonical 三维空间，并量化到十六立方 voxel。同一 voxel 内先聚合像素，再跨有效视图平均，之后进行 LayerNorm 并加入三维位置编码。紧凑的已观测 voxel token 通过 cross-attention 注入 Affostruction 官方 SparseStructureFlowModel。该模型有十二个 block、hidden 七百六十八，全部参数参与微调，训练目标是条件流速度，推理通过 ODE 采样得到 active support。"""
set_notes(s, note); notes.append((3, "SS 微调算法", note))

# 4 — SLat algorithm
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
title(s, 4, "SLat 微调：像素对齐的可信稀疏残差流", "保留 TRELLIS 原生生成能力，仅在有多视图证据的位置进行受控修正")
flow_node(s, "active voxel", "SS support\n64³ 坐标", 0.55, 1.55, 1.72, GREEN)
arrow(s, 2.39, 1.94, 0.31)
flow_node(s, "多视图投影", "K、world-to-camera\n3D → pixel grid", 2.81, 1.55, 1.88, BLUE)
arrow(s, 4.81, 1.94, 0.31)
flow_node(s, "双频采样", "high DINO semantics\n+ low RGB texture", 5.23, 1.55, 1.91, PURPLE)
arrow(s, 7.26, 1.94, 0.31)
flow_node(s, "可信融合", "c × mask × depth\n× angle agreement", 7.68, 1.55, 1.93, CYAN)
arrow(s, 9.73, 1.94, 0.31)
flow_node(s, "稀疏残差流", "768×12 SLat Flow\n预测 Δv", 10.15, 1.55, 1.85, PURPLE)
arrow(s, 12.10, 1.94, 0.27)
add_text(s, "解码", 12.40, 1.78, 0.48, 0.55, 11.5, ORANGE, True, PP_ALIGN.CENTER)

add_rect(s, 0.55, 3.18, 3.83, 2.98, PANEL, line=CYAN)
add_text(s, "逐 voxel 多视图融合", 0.83, 3.43, 3.28, 0.36, 17, CYAN, True)
add_text(s, "fi = Σv αv [fiᴴ ⊕ fiᴸ] / (Σv αv + ε)", 0.81, 4.02, 3.30, 0.48, 14.2, WHITE, True, PP_ALIGN.CENTER)
add_text(s, "αv = confidence · visibility ·\ndepth agreement · angle agreement", 0.86, 4.65, 3.18, 0.72, 12.8, WHITE, False, PP_ALIGN.CENTER)
add_text(s, "输出：1024-D condition + valid mask", 0.88, 5.50, 3.12, 0.30, 11.8, MUTED, True, PP_ALIGN.CENTER)

add_rect(s, 4.65, 3.18, 4.02, 2.98, PANEL, line=PURPLE)
add_text(s, "独立 Affostruction Sparse Flow", 4.93, 3.43, 3.46, 0.36, 17, PURPLE, True)
add_rich_lines(s, [
    "SLatFlowModel：64³ sparse support",
    "8 latent channel、hidden 768",
    "12 blocks / 12 heads / patch size 2",
    "只训练 residual flow 与 conditioner",
], 5.01, 4.02, 3.22, 1.60, 12.8, WHITE, True, 9)

add_rect(s, 8.94, 3.18, 3.83, 2.98, PANEL, line=GREEN)
add_text(s, "冻结先验与受控注入", 9.22, 3.43, 3.28, 0.36, 17, GREEN, True)
add_text(s, "vfinal = vfrozen + c · L · tanh(Δvraw / L)", 9.17, 4.02, 3.37, 0.62, 13.2, WHITE, True, PP_ALIGN.CENTER)
add_rich_lines(s, [
    "TRELLIS native SLat flow 冻结",
    "残差幅值随 base velocity RMS 缩放",
    "低可信区域自动回归原生先验",
], 9.22, 4.78, 3.18, 1.05, 12.3, WHITE, True, 7)
footer(s)
note = """SLat 阶段以 SS 给出的六十四立方 active voxel 为稀疏支撑。系统将每个 voxel 投影回所有有效相机，在对应像素同时采样高层 DINO 语义和浅层卷积提取的 RGB 纹理。每个视图的权重由 VGGT confidence、可见性、前景、深度一致性和入射角一致性共同决定，融合结果投影为一千零二十四维条件 token。一个独立的 Affostruction SLatFlowModel 根据这些条件预测速度残差。TRELLIS 原生 SLat flow 完全冻结，最终修正通过 tanh 限幅并乘以 voxel confidence，所以观测充分的位置得到几何和纹理校正，证据不足的位置保留 TRELLIS 先验。"""
set_notes(s, note); notes.append((4, "SLat 微调算法", note))

# 5 — Training objective and gradient boundary
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
title(s, 5, "训练目标与梯度边界", "把结构生成、先验保持、可信残差分成可审计的优化路径")

add_rect(s, 0.55, 1.47, 5.95, 4.82, PANEL, line=GREEN)
add_text(s, "SS：官方 Conditional Flow Matching", 0.83, 1.72, 5.40, 0.40, 18, GREEN, True)
add_rect(s, 0.86, 2.35, 5.33, 0.75, PANEL_2, line=GREEN)
add_text(s, "xt = (1−t)x0 + [σmin+(1−σmin)t]ε", 1.05, 2.49, 4.95, 0.39, 16, WHITE, True, PP_ALIGN.CENTER)
add_rect(s, 0.86, 3.24, 5.33, 0.75, PANEL_2, line=GREEN)
add_text(s, "v* = (1−σmin)ε − x0    ·    LSS = ‖vθ−v*‖²", 1.01, 3.38, 5.05, 0.39, 15.5, WHITE, True, PP_ALIGN.CENTER)
add_rich_lines(s, [
    "训练：官方 SS Flow 全部参数",
    "冻结：DINO 图像编码器与条件构造",
    "优化：AdamW + EMA + CFG dropout",
], 0.94, 4.27, 5.12, 1.35, 13.5, WHITE, True, 9)
add_rect(s, 0.88, 5.68, 5.27, 0.37, RGBColor(27, 63, 55), line=GREEN)
add_text(s, "单一速度目标保持生成流训练的一致性", 1.02, 5.70, 4.97, 0.29, 12.5, GREEN, True, PP_ALIGN.CENTER)

add_rect(s, 6.78, 1.47, 5.98, 4.82, PANEL, line=PURPLE)
add_text(s, "SLat：冻结先验 + 受控残差", 7.06, 1.72, 5.42, 0.40, 18, PURPLE, True)
add_rect(s, 7.08, 2.35, 5.37, 0.75, PANEL_2, line=PURPLE)
add_text(s, "vfinal = vfrozen + c · L · tanh(Δvraw / L)", 7.25, 2.49, 5.02, 0.39, 15.5, WHITE, True, PP_ALIGN.CENTER)
add_rect(s, 7.08, 3.24, 5.37, 0.75, PANEL_2, line=PURPLE)
add_text(s, "LSLat = Lflow + λx0Lendpoint + λpLprior + Ldecoded", 7.19, 3.38, 5.16, 0.39, 14.5, WHITE, True, PP_ALIGN.CENTER)
add_rich_lines(s, [
    "训练：Affostruction sparse residual flow + conditioner",
    "冻结：TRELLIS native SLat flow / decoder teacher",
    "监督：velocity、端点、低置信先验、解码资产",
], 7.14, 4.27, 5.10, 1.35, 13.1, WHITE, True, 9)
add_rect(s, 7.10, 5.68, 5.31, 0.37, RGBColor(45, 35, 68), line=PURPLE)
add_text(s, "BF16 前向 · FP32 主权重 · DDP · checkpointing", 7.21, 5.70, 5.10, 0.29, 12.0, PURPLE, True, PP_ALIGN.CENTER)
footer(s)
note = """这一页说明梯度到底更新谁。SS 的训练路径与 Affostruction 官方条件流一致：从 clean latent 和高斯噪声构造 xt，目标速度为一减 sigma-min 乘噪声再减 x0，损失是预测速度和目标速度的均方误差。DINO 和条件构造冻结，SS Flow 全参数更新。SLat 以冻结 TRELLIS 的速度为基线，只训练独立稀疏残差流和多视图 conditioner。最终速度等于基础速度加置信度调制后的受限残差。损失同时约束 flow velocity、重建端点、低置信区域的先验保持，以及解码后的资产质量。数值实现使用混合精度前向、FP32 主权重、DDP 和梯度检查点。"""
set_notes(s, note); notes.append((5, "训练目标与梯度边界", note))

# 6 — Inference closure and methodology
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
title(s, 6, "完整推理闭环与总体思路", "显式几何负责可观测区域，生成先验负责稀疏视图下的不确定区域")

flow_node(s, "多视图输入", "精确 frame ID\n缺帧不重排", 0.55, 1.55, 1.82, CYAN)
arrow(s, 2.48, 1.94, 0.32)
flow_node(s, "SS ODE", "CFG + EMA\n生成 support", 2.91, 1.55, 1.70, GREEN)
arrow(s, 4.72, 1.94, 0.32)
flow_node(s, "坐标投影", "active voxel\n→ 所有视图", 5.15, 1.55, 1.74, BLUE)
arrow(s, 7.00, 1.94, 0.32)
flow_node(s, "SLat ODE", "native base\n+ gated residual", 7.43, 1.55, 1.84, PURPLE)
arrow(s, 9.38, 1.94, 0.32)
flow_node(s, "资产解码", "Gaussian + mesh\nheld-out render", 9.81, 1.55, 2.05, ORANGE)

add_rect(s, 0.55, 3.18, 6.02, 2.88, PANEL)
add_text(s, "总体方法：几何条件化生成", 0.83, 3.42, 5.46, 0.38, 17, GREEN, True)
add_rich_lines(s, [
    "SS 把 RGBD 证据转为三维稀疏结构条件",
    "SLat 把 active voxel 转回像素对齐特征",
    "高层语义负责类别与部件，低层 RGB 负责局部纹理",
    "confidence 决定几何观测与生成先验的控制权",
], 0.88, 3.95, 5.28, 1.70, 13.2, WHITE, True, 10)

add_rect(s, 6.81, 3.18, 5.95, 2.88, PANEL, line=ORANGE)
add_text(s, "工程与训练闭环", 7.09, 3.42, 5.40, 0.38, 17, ORANGE, True)
add_rich_lines(s, [
    "frame ID、图像、深度与相机保持一一对应",
    "DDP 利用多 GPU；BF16 前向降低显存与通信成本",
    "FP32 主权重、EMA 与梯度检查点保证稳定训练",
    "训练、验证与推理共享相同坐标系和条件契约",
], 7.13, 3.95, 5.22, 1.70, 13.0, WHITE, True, 10)
add_rect(s, 0.55, 6.28, 12.21, 0.55, RGBColor(23, 57, 61), line=CYAN)
add_text(s, "最终目标：观测区域由 VGGT 精确校正，未观测区域保留 TRELLIS 生成先验", 0.78, 6.36, 11.76, 0.35, 15, CYAN, True, PP_ALIGN.CENTER)
footer(s)
note = """完整推理首先读取严格按 frame ID 对齐的多视图图像、mask、深度和相机。SS 条件流通过 CFG 和 EMA 权重采样 active support。随后将 active voxel 投影到所有有效视图，构造逐 voxel 的 SLat condition。SLat ODE 在每个采样步保留冻结的 TRELLIS base velocity，并叠加置信度门控的残差，最后解码成 Gaussian 和 mesh。总体思路是几何条件化生成：VGGT 为可观测区域提供确定性约束，TRELLIS 为遮挡和未观测区域提供生成先验，高层语义与低层纹理共同决定结构化潜变量，置信度负责在两者之间动态分配控制权。"""
set_notes(s, note); notes.append((6, "推理闭环与总体思路", note))

OUT_DIR.mkdir(parents=True, exist_ok=True)
prs.save(PPT_PATH)

md = ["# 当前 SS / SLat Affostruction 微调架构——演讲稿", ""]
for idx, heading, body in notes:
    md.extend([f"## 第 {idx} 页：{heading}", "", body, ""])
SCRIPT_PATH.write_text("\n".join(md), encoding="utf-8")
print(PPT_PATH)
print(SCRIPT_PATH)
