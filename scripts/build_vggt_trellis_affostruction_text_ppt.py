from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports"
PPT = OUT / "VGGT_TRELLIS_Affostruction架构讲解_4页_文字详解版.pptx"
NOTES = OUT / "VGGT_TRELLIS_Affostruction架构讲解_4页_文字详解版_演讲稿.md"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
BG = RGBColor(8, 17, 30)
PANEL = RGBColor(18, 31, 50)
PANEL2 = RGBColor(24, 40, 62)
WHITE = RGBColor(241, 246, 251)
MUTED = RGBColor(158, 175, 194)
CYAN = RGBColor(54, 207, 220)
BLUE = RGBColor(76, 143, 255)
GREEN = RGBColor(57, 203, 132)
PURPLE = RGBColor(168, 113, 255)
ORANGE = RGBColor(246, 168, 55)


def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def text(slide, value, x, y, w, h, size=14, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margins=0.07):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = Inches(margins); frame.margin_right = Inches(margins)
    frame.margin_top = Inches(margins); frame.margin_bottom = Inches(margins)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value
    r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return shape


def paragraph_box(slide, heading, body, x, y, w, h, color=CYAN, body_size=13.0):
    rect(slide, x, y, w, h, PANEL, line=color)
    text(slide, heading, x + 0.22, y + 0.18, w - 0.44, 0.34, 16, color, True)
    text(slide, body, x + 0.22, y + 0.62, w - 0.44, h - 0.79, body_size, WHITE,
         False, PP_ALIGN.JUSTIFY, MSO_ANCHOR.TOP, 0.02)


def page_title(slide, number, heading, subtitle):
    rect(slide, 0, 0, 0.13, H, CYAN, radius=False)
    text(slide, f"0{number}", 0.45, 0.27, 0.53, 0.35, 12.5, CYAN, True)
    text(slide, heading, 1.02, 0.19, 11.75, 0.56, 25, WHITE, True)
    text(slide, subtitle, 1.04, 0.75, 11.55, 0.29, 10.3, MUTED)
    rect(slide, 0.45, 1.10, 12.34, 0.025, CYAN, radius=False)


def footer(slide):
    text(slide, "VGGT × TRELLIS × Affostruction", 0.47, 7.15, 5.2, 0.18, 8.5, MUTED)


def formula(slide, value, x, y, w, h, color):
    rect(slide, x, y, w, h, PANEL2, line=color)
    text(slide, value, x + 0.12, y + 0.08, w - 0.24, h - 0.16, 14.2, WHITE, True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def flow_chip(slide, heading, x, y, w, color):
    rect(slide, x, y, w, 0.65, PANEL2, line=color)
    text(slide, heading, x + 0.05, y + 0.11, w - 0.10, 0.40, 11.8, color, True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def arrow(slide, x, y, color=CYAN):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(0.31), Inches(0.25))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.color.rgb = color


def notes(slide, value):
    frame = slide.notes_slide.notes_text_frame
    frame.clear(); frame.paragraphs[0].text = value


prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]
scripts = []


# 1. Overall methodology
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 1, "总体方法：由全局图像条件转向三维空间条件", "将 Affostruction 的空间条件化生成思想用于稀疏多视图 VGGT + TRELLIS 重建")

flow_chip(s, "稀疏多视图", 0.62, 1.43, 1.72, CYAN); arrow(s, 2.44, 1.63)
flow_chip(s, "VGGT 几何", 2.86, 1.43, 1.72, BLUE); arrow(s, 4.68, 1.63)
flow_chip(s, "SS 空间条件流", 5.10, 1.43, 1.94, GREEN); arrow(s, 7.14, 1.63)
flow_chip(s, "SLat 可信残差流", 7.56, 1.43, 2.08, PURPLE); arrow(s, 9.74, 1.63)
flow_chip(s, "TRELLIS 解码", 10.16, 1.43, 2.20, ORANGE)

paragraph_box(
    s, "核心思路",
    "原始 VGGT 能够从稀疏多视图中估计相机、深度、点图和置信度，TRELLIS 则拥有强大的三维生成先验，但二者之间缺少逐三维位置的显式信息通道。当前架构吸收 Affostruction 的关键思想：不再只把整张图像压缩成全局条件，而是利用深度和相机把图像特征附着到三维 voxel，使生成流中的每一个结构位置都能够通过 cross-attention 读取与自身对应的图像证据。",
    0.62, 2.39, 7.57, 2.31, CYAN, 13.2,
)
paragraph_box(
    s, "面向图像重建的适配",
    "项目不使用文本、CLIP 或 affordance prompt，而是把条件改造成 VGGT 几何、DINO 高层语义和 RGB 低层纹理。SS 阶段负责确定“哪里存在结构”；SLat 阶段负责确定“每个结构点具有怎样的局部几何与外观”。观测充分的位置由多视图证据校正，遮挡和未观测区域继续依赖 TRELLIS 的生成能力。",
    8.45, 2.39, 4.31, 2.31, PURPLE, 12.6,
)
rect(s, 0.62, 5.02, 12.14, 1.27, RGBColor(20, 52, 58), line=CYAN)
text(s, "Pixel → Voxel → Pixel-aligned Evidence → Gated Structured Latent", 0.91, 5.19, 11.56, 0.37, 17, CYAN, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, "二维观测先形成三维结构条件，再由三维结构主动回到图像中检索对应证据，构成双向空间闭环。", 1.05, 5.69, 11.28, 0.34, 13.2, WHITE, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
footer(s)
n = """当前架构的核心变化，是把图像条件从全局语义提升为三维空间条件。VGGT 提供相机、深度、点图和置信度，TRELLIS 提供完整三维对象的生成先验。项目在两者之间引入 Affostruction 风格的空间 token 和条件流：SS 把像素证据附着到三维 voxel，SLat 再从 active voxel 回到所有图像中读取对应语义与纹理。这样形成 Pixel 到 Voxel，再回到 Pixel-aligned Evidence 的双向闭环。项目只吸收适合图像三维重建的空间条件化、cross-attention 和 flow matching，不采用文本或 affordance prompt。"""
notes(s, n); scripts.append((1, "总体方法", n))


# 2. SS
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 2, "SS 微调：RGBD 体素条件驱动的官方生成流", "VGGT 定义几何位置，DINO 提供局部语义，Affostruction SS Flow 学习结构生成速度场")

paragraph_box(
    s, "RGBD 空间条件如何形成",
    "每个输入视图先由冻结的 DINOv2 ViT-L/14 提取 1024 维 patch feature。对应像素利用深度、相机内参 K 和 camera-to-world 变换恢复到 canonical 三维空间，再量化到 16³ voxel。同一 voxel 内先计算平均二维采样位置，通过 grid_sample 读取 DINO 特征；随后只在真实有效视图之间进行平均。聚合特征经过 LayerNorm，并加入三维位置编码，使 token 同时携带“这里是什么”以及“它位于三维空间哪里”。只有被至少一个视图观测到的 voxel 才进入紧凑 condition 序列，缺失或 padding 视图由 valid mask 排除。",
    0.62, 1.42, 7.52, 2.55, BLUE, 12.65,
)
paragraph_box(
    s, "模型如何利用这些条件",
    "SS 使用 Affostruction 官方 SparseStructureFlowModel，而不是额外叠加小型 Adapter。模型工作在 16³、8 通道的 SS latent 上，hidden width 为 768，包含 12 个 Transformer block 和 12 个 attention head。带噪 latent 通过 self-attention 建立结构内部关系，RGBD voxel token 则通过每层 cross-attention 注入；时间嵌入调制不同噪声阶段的去噪行为。DINO 和条件构造被冻结，SS Flow 的全部参数参与微调，因此模型能够整体适配新的三维条件空间。",
    0.62, 4.18, 7.52, 2.25, GREEN, 12.65,
)

rect(s, 8.43, 1.42, 4.33, 1.33, PANEL, line=GREEN)
text(s, "官方拓扑", 8.68, 1.62, 1.18, 0.31, 15.5, GREEN, True)
text(s, "16³ · in/out 8 · hidden 768\n12 blocks · 12 heads · APE · Q/K RMS norm", 8.68, 2.00, 3.77, 0.55, 12.3, WHITE)

rect(s, 8.43, 2.98, 4.33, 2.21, PANEL, line=GREEN)
text(s, "Conditional Flow Matching", 8.68, 3.18, 3.83, 0.33, 15.5, GREEN, True)
formula(s, "xₜ=(1−t)x₀+[σmin+(1−σmin)t]ε", 8.72, 3.67, 3.74, 0.53, GREEN)
formula(s, "v*=(1−σmin)ε−x₀", 8.72, 4.32, 3.74, 0.53, GREEN)

rect(s, 8.43, 5.42, 4.33, 1.01, RGBColor(26, 62, 53), line=GREEN)
text(s, "LSS = ‖vθ(xₜ,t,C)−v*‖²", 8.68, 5.56, 3.83, 0.34, 16, GREEN, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, "ODE + CFG + EMA → active voxel support", 8.64, 5.94, 3.91, 0.24, 10.8, WHITE, False, PP_ALIGN.CENTER)
footer(s)
n = """SS 阶段先把多视图 RGBD 变成与三维位置对应的条件。冻结 DINO 提取空间 patch 特征，深度与相机把像素反投影到 canonical 空间并量化到十六立方 voxel；同一 voxel 内通过平均二维位置进行特征采样，再跨有效视图平均。LayerNorm 和三维位置编码使 token 同时具有语义与位置。官方 Affostruction SparseStructureFlowModel 有十二层、宽度七百六十八，全部参数微调，条件 token 通过 cross-attention 进入每个生成 block。训练使用 Conditional Flow Matching，仅拟合正确的 latent velocity；推理利用 EMA 权重、CFG 和 ODE 积分得到 active support。"""
notes(s, n); scripts.append((2, "SS 微调", n))


# 3. SLat
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 3, "SLat 微调：像素对齐的可信稀疏残差流", "冻结 TRELLIS 原生 SLat 先验，只让图像条件在有可靠证据的 voxel 上产生受控修正")

paragraph_box(
    s, "从 active voxel 回到图像检索证据",
    "SLat 以 SS 输出的 64³ active voxel 为稀疏支撑。系统使用每个视图的 K 和 world-to-camera 变换，将 voxel 投影为二维 sampling grid，并在相同位置同时采样两类信息：高层 DINO/VGGT 特征描述类别、部件和语义关系；浅层 RGB CNN 保留颜色、边缘与局部纹理。每个视图还会采样 foreground mask、aligned depth、VGGT confidence 与 point-map normal，从而判断该观测是否落在前景、是否被遮挡，以及投影深度和观察角度是否与当前三维点一致。",
    0.62, 1.42, 7.45, 2.47, PURPLE, 12.7,
)
paragraph_box(
    s, "为什么采用冻结先验上的残差学习",
    "融合后的 1024 维空间 token 通过 cross-attention 输入独立的 Affostruction SLatFlowModel，由它预测相对 TRELLIS 基础速度的修正量，而不是重新生成完整速度。TRELLIS 原生 SLat flow 保持冻结，继续负责遮挡和未观测区域的生成完整性。残差幅值按照 base velocity 的 RMS 自适应缩放，并经过 tanh 限幅；最终再乘 voxel confidence。因此只有多视图证据充分的位置能够明显修改 latent，低可信位置会自动退回原始 TRELLIS 先验。",
    0.62, 4.12, 7.45, 2.31, GREEN, 12.7,
)

rect(s, 8.36, 1.42, 4.40, 1.62, PANEL, line=CYAN)
text(s, "多视图融合", 8.62, 1.63, 3.86, 0.32, 15.5, CYAN, True)
formula(s, "αᵢᵛ=cᵢᵛ·mᵢᵛ·wdepth·wangle", 8.65, 2.05, 3.82, 0.47, CYAN)
text(s, "fᵢ=Σᵥαᵢᵛ[fᴴ⊕fᴸ]/(Σᵥαᵢᵛ+ε)", 8.61, 2.58, 3.89, 0.30, 11.5, WHITE, True, PP_ALIGN.CENTER)

rect(s, 8.36, 3.28, 4.40, 1.25, PANEL, line=PURPLE)
text(s, "稀疏残差模型", 8.62, 3.49, 3.86, 0.32, 15.5, PURPLE, True)
text(s, "64³ sparse support · in/out 8 · hidden 768\n12 blocks · 12 heads · patch size 2", 8.62, 3.89, 3.86, 0.49, 12.0, WHITE)

rect(s, 8.36, 4.78, 4.40, 1.65, PANEL, line=GREEN)
text(s, "置信度门控修正", 8.62, 4.99, 3.86, 0.32, 15.5, GREEN, True)
formula(s, "L=λres·RMS(vbase)", 8.65, 5.39, 3.82, 0.42, GREEN)
text(s, "vfinal=vfrozen+c·L·tanh(Δvraw/L)", 8.55, 5.90, 4.02, 0.31, 12.0, GREEN, True, PP_ALIGN.CENTER)
footer(s)
n = """SLat 阶段把 SS active voxel 投影回所有输入视图，在同一像素位置采样高层语义和低层纹理，并结合前景、遮挡、深度、法向角度和 VGGT confidence 计算逐 voxel、逐视图权重。融合后的空间 token 输入独立 Affostruction SLatFlowModel。这里没有覆盖 TRELLIS 原生 SLat flow，而是冻结它作为 base velocity。新模型只学习 residual，残差按基础速度 RMS 缩放，通过 tanh 限幅并乘 voxel confidence。于是图像证据只在可靠区域发挥作用，而遮挡或未观测区域继续保留 TRELLIS 的生成先验。"""
notes(s, n); scripts.append((3, "SLat 微调", n))


# 4. Closed loop
s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
page_title(s, 4, "训练与推理闭环：观测约束与生成先验协同", "SS 学习结构速度场，SLat 学习条件残差；训练、采样与解码共享同一空间契约")

paragraph_box(
    s, "训练阶段",
    "SS 训练时冻结 DINO 和 RGBD conditioner，只更新 Affostruction SS Flow，并以 Conditional Flow Matching 的速度 MSE 为主目标。SLat 训练时冻结 TRELLIS native SLat flow，把它的输出作为 base teacher；可训练部分只有 sparse residual flow 与像素对齐 conditioner。SLat 总目标同时约束最终 velocity、由速度反推的 clean endpoint、低置信区域的 prior preservation，以及经 TRELLIS decoder 得到的 RGB、mask、depth、feature consistency 与 surface geometry。",
    0.62, 1.42, 6.05, 2.58, GREEN, 12.8,
)
paragraph_box(
    s, "推理阶段",
    "输入端首先按照真实 frame ID 同步读取 RGB、mask、depth、K 与 T，避免缺失编号造成视图和相机错位。SS 使用 EMA、CFG 和 ODE 从多视图三维条件生成 active support。SLat 在每个 ODE 步先计算冻结 TRELLIS base velocity，再构造 active voxel 的多视图像素条件并叠加 gated residual。最终 Structured Latent 被送入 TRELLIS Gaussian 或 mesh decoder，生成可用于新视角渲染与几何评估的三维资产。",
    6.93, 1.42, 5.83, 2.58, PURPLE, 12.8,
)

rect(s, 0.62, 4.28, 3.82, 1.69, PANEL, line=GREEN)
text(s, "SS 梯度边界", 0.88, 4.50, 3.30, 0.31, 15.5, GREEN, True)
text(s, "训练：SS Flow 全参数\n冻结：DINO + RGBD conditioner\n目标：velocity CFM-MSE + EMA", 0.88, 4.94, 3.30, 0.77, 12.4, WHITE)

rect(s, 4.76, 4.28, 4.04, 1.69, PANEL, line=PURPLE)
text(s, "SLat 梯度边界", 5.02, 4.50, 3.52, 0.31, 15.5, PURPLE, True)
text(s, "训练：residual flow + conditioner\n冻结：TRELLIS native SLat flow\n目标：flow + endpoint + prior + decoded", 5.02, 4.94, 3.52, 0.77, 12.2, WHITE)

rect(s, 9.12, 4.28, 3.64, 1.69, PANEL, line=CYAN)
text(s, "工程契约", 9.38, 4.50, 3.12, 0.31, 15.5, CYAN, True)
text(s, "统一 canonical frame / camera convention\nDDP 多 GPU · BF16 前向 · FP32 主权重\ngradient checkpointing 控制激活显存", 9.38, 4.94, 3.12, 0.77, 11.7, WHITE)

rect(s, 0.62, 6.23, 12.14, 0.58, RGBColor(20, 52, 58), line=CYAN)
text(s, "最终原则：可观测区域由 VGGT 精确校正，未观测区域由 TRELLIS 保持合理生成。", 0.91, 6.33, 11.56, 0.33, 15.2, CYAN, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
footer(s)
n = """训练阶段严格划分梯度边界。SS 只更新官方 Affostruction Flow，DINO 和 RGBD 条件冻结；SLat 只更新 residual flow 和 conditioner，TRELLIS 原生 SLat flow 冻结。SLat 除了速度残差，还通过 endpoint、prior preservation 和解码资产损失约束最终三维结果。推理阶段按 frame ID 同步全部输入，SS 先生成 active support，SLat 再在每个 ODE 步把 frozen base 与 gated residual 相加，最后交给 TRELLIS decoder。训练与推理必须共享 canonical frame、相机约定和 mask 契约。整套设计的最终原则，是让 VGGT 精确约束可观测区域，让 TRELLIS 完成不确定区域。"""
notes(s, n); scripts.append((4, "训练与推理闭环", n))


OUT.mkdir(parents=True, exist_ok=True)
prs.save(PPT)
md = ["# VGGT + TRELLIS + Affostruction 架构讲解——文字详解版演讲稿", ""]
for index, heading, body in scripts:
    md.extend([f"## 第 {index} 页：{heading}", "", body, ""])
NOTES.write_text("\n".join(md), encoding="utf-8")
print(PPT)
print(NOTES)
